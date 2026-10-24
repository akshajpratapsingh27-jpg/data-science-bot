"""Report generation for dsbot - narrative reports from data using LLM."""

import json
import re
import shutil
import subprocess
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.shared import Inches
from pydantic import BaseModel

from dsbot.approval import ApprovalRequest, ApprovalType, get_approval_queue
from dsbot.audit import ActionType, get_audit_log, log
from dsbot.data.analysis import Facts, analyze
from dsbot.data.clean import ExtractionResult
from dsbot.data.pipeline import PipelineResult
from dsbot.llm import LLM, get_llm


class ReportConfig(BaseModel):
    """Configuration for report generation."""
    title: str = "Data Analysis Report"
    sections: list[str] = ["summary", "key_findings", "statistics", "recommendations"]
    max_length_words: int = 500
    include_charts: bool = False


class DataReport(BaseModel):
    """A generated data report."""
    title: str
    source_name: str
    content: str  # Markdown content
    metadata: dict[str, Any] = {}
    generated_at: str  # ISO format
    approval_id: str | None = None

    model_config = {"arbitrary_types_allowed": True}


class ReportGenerator:
    """Generates narrative reports from data using LLM."""

    def __init__(self, llm: LLM, config: ReportConfig | None = None):
        self.llm = llm
        self.config = config or ReportConfig()

    def generate(self, result: ExtractionResult) -> DataReport:
        """Generate a report from cleaned data."""
        # Compute statistics using pandas (never LLM)
        stats = self._compute_statistics(result.dataframe)

        # Build prompt with actual data statistics
        prompt = self._build_prompt(result, stats)

        response = self.llm.complete(self._system_prompt(), prompt)

        # Create report
        from datetime import datetime
        report = DataReport(
            title=self.config.title,
            source_name=result.source_name,
            content=response,
            metadata={
                "rows": len(result.dataframe),
                "columns": list(result.dataframe.columns),
                "statistics": stats,
                "cleaning_steps": result.metadata.get("cleaning_steps_applied", []),
            },
            generated_at=datetime.utcnow().isoformat(),
        )

        return report

    def _compute_statistics(self, df: pd.DataFrame) -> dict[str, Any]:
        """Compute descriptive statistics using pandas."""
        stats = {
            "shape": df.shape,
            "dtypes": df.dtypes.astype(str).to_dict(),
            "missing_values": df.isnull().sum().to_dict(),
            "numeric_summary": {},
            "categorical_summary": {},
        }

        # Numeric columns
        numeric_cols = df.select_dtypes(include=["number"]).columns
        if len(numeric_cols) > 0:
            stats["numeric_summary"] = df[numeric_cols].describe().to_dict()

        # Categorical columns
        cat_cols = df.select_dtypes(include=["object", "category"]).columns
        for col in cat_cols:
            stats["categorical_summary"][col] = {
                "unique_count": df[col].nunique(),
                "top_values": df[col].value_counts().head(5).to_dict(),
            }

        return stats

    def _system_prompt(self) -> str:
        return """You are a data scientist writing a professional report.
Write a clear, concise report based on the provided data statistics.

Guidelines:
- Use the ACTUAL statistics provided - never invent numbers
- Write in professional business language
- Structure with clear headings
- Focus on insights, not just listing numbers
- Keep it under the specified word limit
- Use markdown formatting

Return ONLY the report content as markdown."""

    def _build_prompt(self, result: ExtractionResult, stats: dict) -> str:
        sections_str = ", ".join(self.config.sections)
        return f"""Generate a report for: {result.source_name}

Data shape: {stats['shape'][0]} rows, {stats['shape'][1]} columns
Columns: {list(result.dataframe.columns)}
Dtypes: {stats['dtypes']}
Missing values: {stats['missing_values']}

Numeric summary:
{json.dumps(stats['numeric_summary'], indent=2, default=str)}

Categorical summary:
{json.dumps(stats['categorical_summary'], indent=2, default=str)}

Cleaning steps applied: {result.metadata.get('cleaning_steps_applied', [])}

Report sections to include: {sections_str}
Max length: ~{self.config.max_length_words} words

Write the report now."""


def queue_report_for_approval(
    report: DataReport,
    approval_queue=None,
    audit_log=None,
) -> ApprovalRequest:
    """Queue a generated report for human approval before release."""
    if approval_queue is None:
        approval_queue = get_approval_queue()
    if audit_log is None:
        audit_log = get_audit_log()

    request = ApprovalRequest(
        type=ApprovalType.DATA_REPORT_RELEASE,
        payload=report.model_dump(),
        metadata={
            "source_name": report.source_name,
            "title": report.title,
        },
    )

    approval_queue.add(request)
    audit_log.log_action(
        ActionType.DATA_REPORT_GENERATE,
        actor="bot",
        details={"approval_id": request.id, "source": report.source_name, "title": report.title},
        correlation_id=request.id,
    )

    # Update report with approval ID
    report.approval_id = request.id
    return request


def save_report(report: DataReport, output_path: Path | str) -> Path:
    """Save report to markdown file."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(report.content, encoding="utf-8")
    return output_path


NUMBER_PATTERN = re.compile(r"(?<![A-Za-z_])[-+]?(?:\d[\d,]*)(?:\.\d+)?%?")


def _number_tokens(text: str) -> list[str]:
    return NUMBER_PATTERN.findall(text)


def _as_number(token: str) -> float:
    return float(token.rstrip("%").replace(",", ""))


def _facts_json(facts: Facts) -> str:
    return json.dumps(facts, default=str, sort_keys=True, indent=2)


def _narrative_numbers_are_supported(narrative: str, facts: Facts) -> bool:
    """Ensure every number in generated prose exists in the computed facts."""

    fact_tokens = _number_tokens(_facts_json(facts))
    allowed = [_as_number(token) for token in fact_tokens]
    for token in _number_tokens(narrative):
        try:
            candidate = _as_number(token)
        except ValueError:
            return False
        if not any(round(candidate, 2) == round(value, 2) for value in allowed):
            return False
    return True


def _format_fact_value(value: Any) -> str:
    if isinstance(value, float):
        return f"{value:.2f}".rstrip("0").rstrip(".")
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, default=str, sort_keys=True)
    return str(value)


def _fallback_narrative(facts: Facts) -> str:
    """Build a deterministic narrative without adding unsupported numbers."""

    lines = ["The following findings were computed directly from the cleaned data:"]
    for fact in facts:
        template = str(fact.get("statement_template", "")).strip()
        if not template:
            continue
        lines.append(f"- {template.replace('{value}', _format_fact_value(fact.get('value')))}")
    return "\n".join(lines)


def _narrative_from_llm(facts: Facts, llm: LLM) -> str:
    facts_json = _facts_json(facts)
    system = (
        "You are a careful data analyst. Write concise professional narrative only from the "
        "Facts JSON supplied by the user. Do not invent facts, numbers, dates, percentages, "
        "or commitments. Return prose only."
    )
    prompt = (
        "Facts JSON follows. Use no information outside this JSON and do not add unsupported "
        "numbers.\n\n"
        f"{facts_json}"
    )
    first = llm.complete(system, prompt)
    if _narrative_numbers_are_supported(first, facts):
        return first

    retry_prompt = (
        "The previous narrative contained unsupported numeric content. Rewrite it using only "
        "the Facts JSON below. If a value is not present, omit it. Return prose only.\n\n"
        f"{facts_json}"
    )
    second = llm.complete(system, retry_prompt)
    if _narrative_numbers_are_supported(second, facts):
        return second
    return _fallback_narrative(facts)


def _profile_value(profile: dict[str, Any], key: str, default: Any = 0) -> Any:
    return profile.get(key, default)


def _add_data_quality_section(document: Any, result: PipelineResult) -> None:
    document.add_heading("Data Quality", level=1)
    before = result.profile_before
    after = result.profile_after
    quality_lines = [
        f"Rows before cleaning: {_profile_value(before, 'row_count')}",
        f"Rows after cleaning: {_profile_value(after, 'row_count')}",
        f"Duplicate rows before cleaning: {_profile_value(before, 'duplicate_count')}",
        f"Duplicate rows after cleaning: {_profile_value(after, 'duplicate_count')}",
        f"Validation status: {'valid' if result.validation_report.get('valid') else 'invalid'}",
    ]
    for line in quality_lines:
        document.add_paragraph(line, style="List Bullet")
    document.add_paragraph("Cleaning change log:")
    for change in result.change_log:
        document.add_paragraph(
            f"{change.step}: {change.rows_affected} rows affected",
            style="List Bullet 2",
        )


def _chart_paths() -> list[Path]:
    return sorted(Path("outputs/charts").glob("*.png"))


def _export_pdf(docx_path: Path) -> None:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        return
    try:
        subprocess.run(
            [
                executable,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(docx_path.parent),
                str(docx_path),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=120,
        )
    except (OSError, subprocess.SubprocessError):
        return


def build_report(
    result: PipelineResult,
    facts: Facts,
    out_path: str | Path,
    llm: LLM | None = None,
) -> Path:
    """Create a DOCX report whose generated prose is guarded against invented numbers."""

    output_path = Path(out_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    active_llm = llm or get_llm()
    narrative = _narrative_from_llm(facts, active_llm)

    charts = _chart_paths()
    if not charts:
        analyze(result.clean_df)
        charts = _chart_paths()

    document = Document()
    document.add_heading("Data Analysis Report", level=0)
    document.add_paragraph(f"Generated: {datetime.now(UTC).isoformat()}")

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(narrative)

    _add_data_quality_section(document, result)

    document.add_heading("Key Findings", level=1)
    document.add_paragraph(narrative)
    for chart in charts[:3]:
        document.add_picture(str(chart), width=Inches(6.2))
        document.add_paragraph(chart.name)

    document.add_heading("Method", level=1)
    document.add_paragraph(
        "The source was extracted, profiled, cleaned, validated, and analyzed with pandas. "
        "Charts were generated with matplotlib. Narrative text was constrained to the "
        "computed Facts JSON."
    )

    document.add_heading("Appendix", level=1)
    document.add_paragraph("Computed Facts JSON:")
    document.add_paragraph(_facts_json(facts))
    document.save(str(output_path))
    _export_pdf(output_path)
    log("bot", "data_report_generate", path=str(output_path), facts=len(facts))
    return output_path


def build_markdown_report(
    result: PipelineResult,
    llm: LLM | None = None,
    charts_dir: str | Path = "outputs/charts",
) -> str:
    """Build a Markdown report whose numeric claims are grounded in pandas facts."""

    facts = analyze(result.clean_df, charts_dir=charts_dir)
    narrative = _narrative_from_llm(facts, llm or get_llm())
    changes = [
        f"- {change.step}: {change.rows_affected} rows affected"
        for change in result.change_log
    ] or ["- No cleaning changes were required."]
    return "\n".join(
        [
            "# Data Analysis Report",
            "",
            "## Summary",
            narrative,
            "",
            "## Data quality",
            f"- Rows before cleaning: {len(result.raw_df)}",
            f"- Rows after cleaning: {len(result.clean_df)}",
            f"- Validation: {'passed' if result.validation_report.get('valid') else 'failed'}",
            "",
            "## Cleaning change log",
            *changes,
            "",
            "## Pandas-computed facts",
            "```json",
            _facts_json(facts),
            "```",
            "",
        ]
    )
