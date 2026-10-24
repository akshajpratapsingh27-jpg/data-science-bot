"""Excel export for dsbot."""

import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from xlsxwriter.utility import xl_col_to_name

from dsbot.audit import log
from dsbot.data.clean import ExtractionResult
from dsbot.data.pipeline import PipelineResult


class ExcelExportConfig(BaseModel):
    """Configuration for Excel export."""
    sheet_name: str = "Data"
    include_index: bool = False
    auto_filter: bool = True
    freeze_header: bool = True
    header_style: dict[str, Any] = {
        "bold": True,
        "fill_color": "4472C4",
        "font_color": "FFFFFF",
        "alignment": "center",
    }
    column_widths: dict[str, int] | None = None
    conditional_formats: list[dict[str, Any]] = []


class ExcelExporter:
    """Exports cleaned data to Excel with formatting."""

    def __init__(self, config: ExcelExportConfig | None = None):
        self.config = config or ExcelExportConfig()

    def export(self, result: ExtractionResult, output_path: Path | str) -> Path:
        """Export DataFrame to Excel with formatting."""
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        df = result.dataframe

        # Create workbook using openpyxl for formatting control
        wb = Workbook()
        ws = wb.active
        ws.title = self.config.sheet_name[:31]  # Excel sheet name limit

        # Write headers
        headers = list(df.columns)
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            self._style_header_cell(cell)

        # Write data
        for row_idx, row in enumerate(df.itertuples(index=self.config.include_index), 2):
            if self.config.include_index:
                # First column is index
                ws.cell(row=row_idx, column=1, value=row.Index)
                data_start = 2
            else:
                data_start = 1

            for col_idx, value in enumerate(row[data_start:], data_start):
                ws.cell(row=row_idx, column=col_idx, value=self._format_cell_value(value))

        # Apply formatting
        self._apply_formatting(ws, len(headers), len(df) + 1)

        wb.save(output_path)
        return output_path

    def _style_header_cell(self, cell) -> None:
        """Apply header styling."""
        style = self.config.header_style
        cell.font = Font(bold=style.get("bold", True), color=style.get("font_color", "FFFFFF"))
        cell.fill = PatternFill(start_color=style.get("fill_color", "4472C4"),
                                end_color=style.get("fill_color", "4472C4"),
                                fill_type="solid")
        cell.alignment = Alignment(horizontal=style.get("alignment", "center"))

    def _format_cell_value(self, value: Any) -> Any:
        """Format cell value for Excel."""
        if pd.isna(value):
            return ""
        if isinstance(value, (pd.Timestamp,)):
            return value.to_pydatetime()
        if isinstance(value, (np.integer,)):
            return int(value)
        if isinstance(value, (np.floating,)):
            return float(value)
        return value

    def _apply_formatting(self, ws, num_cols: int, num_rows: int) -> None:
        """Apply worksheet formatting."""
        # Auto filter
        if self.config.auto_filter and num_rows > 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(num_cols)}{num_rows - 1}"

        # Freeze header row
        if self.config.freeze_header:
            ws.freeze_panes = "A2"

        # Column widths
        if self.config.column_widths:
            for col_name, width in self.config.column_widths.items():
                try:
                    col_idx = (
                        list(ws.iter_cols(min_row=1, max_row=1, values_only=True))[0].index(
                            col_name
                        )
                        + 1
                    )
                    ws.column_dimensions[get_column_letter(col_idx)].width = width
                except (ValueError, IndexError):
                    pass
        else:
            # Auto-width based on content (approximate)
            for col_idx in range(1, num_cols + 1):
                max_length = 0
                for row in ws.iter_rows(
                    min_col=col_idx,
                    max_col=col_idx,
                    max_row=min(num_rows, 50),
                ):
                    for cell in row:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 50)

        # Conditional formatting
        for cf in self.config.conditional_formats:
            self._apply_conditional_format(ws, cf, num_rows, num_cols)

    def _apply_conditional_format(self, ws, cf: dict, num_rows: int, num_cols: int) -> None:
        """Apply conditional formatting rule."""
        from openpyxl.formatting.rule import CellIsRule, ColorScaleRule

        range_str = f"A2:{get_column_letter(num_cols)}{num_rows - 1}"

        if cf.get("type") == "color_scale":
            ws.conditional_formatting.add(range_str, ColorScaleRule(
                start_type='min', start_color=cf.get("start_color", "F8696B"),
                end_type='max', end_color=cf.get("end_color", "63BE7B"),
            ))
        elif cf.get("type") == "cell_is":
            ws.conditional_formatting.add(range_str, CellIsRule(
                operator=cf.get("operator", "greaterThan"),
                formula=[str(cf.get("value", 0))],
                fill=PatternFill(start_color=cf.get("color", "FFC7CE"), fill_type="solid"),
            ))


def export_to_excel(
    result: ExtractionResult,
    output_path: Path | str,
    config: ExcelExportConfig | None = None,
) -> Path:
    """Convenience function to export to Excel."""
    exporter = ExcelExporter(config)
    return exporter.export(result, output_path)


def _xlsx_value(value: Any) -> Any:
    """Convert pandas and numpy values to values xlsxwriter can write."""

    if value is None or pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, default=str, sort_keys=True)
    return value


def _write_dataframe(worksheet: Any, dataframe: pd.DataFrame, formats: dict[str, Any]) -> None:
    """Write a dataframe with filters, widths, and basic type formats."""

    header_format = formats["header"]
    date_format = formats["date"]
    number_format = formats["number"]
    columns = [str(column) for column in dataframe.columns]
    for column_index, column in enumerate(columns):
        worksheet.write(0, column_index, column, header_format)

    for row_index, (_, row) in enumerate(dataframe.iterrows(), start=1):
        for column_index, column in enumerate(columns):
            value = row[column]
            if value is None or pd.isna(value):
                worksheet.write_blank(row_index, column_index, None)
            elif pd.api.types.is_datetime64_any_dtype(dataframe[column]):
                worksheet.write_datetime(
                    row_index,
                    column_index,
                    value.to_pydatetime(),
                    date_format,
                )
            elif (
                pd.api.types.is_numeric_dtype(dataframe[column])
                and not pd.api.types.is_bool_dtype(dataframe[column])
            ):
                worksheet.write_number(row_index, column_index, float(value), number_format)
            elif isinstance(value, (bool, np.bool_)):
                worksheet.write_boolean(row_index, column_index, bool(value))
            else:
                worksheet.write(row_index, column_index, _xlsx_value(value))

    if columns:
        worksheet.freeze_panes(1, 0)
        worksheet.autofilter(0, 0, max(len(dataframe), 1), len(columns) - 1)
        for column_index, column in enumerate(columns):
            values = dataframe[column].map(lambda value: "" if pd.isna(value) else str(value))
            lengths = values.head(100).map(len).tolist()
            longest = max([len(column), *lengths], default=len(column))
            worksheet.set_column(column_index, column_index, min(max(longest + 2, 10), 36))


def _dataframe_column_index(dataframe: pd.DataFrame, column: str) -> int:
    return list(map(str, dataframe.columns)).index(column)


def _write_summary(
    workbook: Any,
    result: PipelineResult,
    formats: dict[str, Any],
) -> None:
    worksheet = workbook.add_worksheet("Summary")
    title_format = formats["title"]
    label_format = formats["label"]
    number_format = formats["number"]
    worksheet.write("A1", "Data Pipeline Summary", title_format)
    worksheet.write("A2", "Generated at")
    worksheet.write_datetime(
        "B2",
        datetime.now(UTC).replace(tzinfo=None),
        formats["date_time"],
    )

    kpis = [
        ("Raw rows", len(result.raw_df)),
        ("Cleaned rows", len(result.clean_df)),
        ("Rows removed", len(result.raw_df) - len(result.clean_df)),
        ("Columns", len(result.clean_df.columns)),
        ("Duplicate rows before", result.profile_before.get("duplicate_count", 0)),
        ("Duplicate rows after", result.profile_after.get("duplicate_count", 0)),
        ("Validation valid", result.validation_report.get("valid", False)),
    ]
    worksheet.write_row("A4", ["KPI", "Value"], formats["header"])
    for row_index, (label, value) in enumerate(kpis, start=4):
        worksheet.write(row_index, 0, label, label_format)
        if isinstance(value, bool):
            worksheet.write_boolean(row_index, 1, value)
        else:
            worksheet.write_number(row_index, 1, float(value), number_format)

    start_row = 4
    start_col = 3
    worksheet.write_row(
        start_row,
        start_col,
        ["Pandas statistics", "Pandas value", "Excel formula"],
        formats["header"],
    )
    numeric_columns = [
        str(column)
        for column in result.clean_df.select_dtypes(include="number").columns
        if not pd.api.types.is_bool_dtype(result.clean_df[column])
    ]
    stat_rows = [
        ("total", "sum", "SUM"),
        ("average", "mean", "AVERAGE"),
        ("minimum", "min", "MIN"),
        ("maximum", "max", "MAX"),
    ]
    row = start_row + 1
    for column in numeric_columns:
        pandas_series = pd.to_numeric(result.clean_df[column], errors="coerce")
        column_index = _dataframe_column_index(result.clean_df, column)
        excel_column = xl_col_to_name(column_index)
        first_data_row = 2
        last_data_row = max(len(result.clean_df) + 1, first_data_row)
        for display_name, pandas_method, excel_function in stat_rows:
            pandas_value = getattr(pandas_series, pandas_method)()
            worksheet.write(row, start_col, f"{column} {display_name}")
            worksheet.write_number(row, start_col + 1, float(pandas_value), number_format)
            formula = (
                f"={excel_function}('Cleaned Data'!"
                f"{excel_column}{first_data_row}:{excel_column}{last_data_row})"
            )
            worksheet.write_formula(row, start_col + 2, formula, number_format, float(pandas_value))
            row += 1

    worksheet.set_column("A:A", 24)
    worksheet.set_column("B:B", 16)
    worksheet.set_column("D:F", 26)


def _write_change_log(workbook: Any, result: PipelineResult, formats: dict[str, Any]) -> None:
    worksheet = workbook.add_worksheet("Change Log")
    headers = ["Step", "Rows Affected", "Columns Affected", "Details"]
    worksheet.write_row(0, 0, headers, formats["header"])
    for row_index, change in enumerate(result.change_log, start=1):
        worksheet.write(row_index, 0, change.step)
        worksheet.write_number(row_index, 1, change.rows_affected, formats["number"])
        worksheet.write(row_index, 2, ", ".join(change.columns_affected))
        worksheet.write(row_index, 3, json.dumps(change.details, default=str, sort_keys=True))
    worksheet.freeze_panes(1, 0)
    worksheet.autofilter(0, 0, max(len(result.change_log), 1), len(headers) - 1)
    worksheet.set_column("A:A", 30)
    worksheet.set_column("B:B", 16)
    worksheet.set_column("C:C", 30)
    worksheet.set_column("D:D", 70)


def _write_dictionary(workbook: Any, result: PipelineResult, formats: dict[str, Any]) -> None:
    worksheet = workbook.add_worksheet("Data Dictionary")
    headers = ["Column", "Dtype", "Description", "Null %"]
    worksheet.write_row(0, 0, headers, formats["header"])
    null_percent = result.profile_after.get("null_percent", {})
    for row_index, column in enumerate(result.clean_df.columns, start=1):
        name = str(column)
        worksheet.write(row_index, 0, name)
        worksheet.write(row_index, 1, str(result.clean_df[column].dtype))
        worksheet.write(row_index, 2, "[PLACEHOLDER]")
        worksheet.write_number(row_index, 3, float(null_percent.get(name, 0.0)), formats["percent"])
    worksheet.freeze_panes(1, 0)
    worksheet.autofilter(0, 0, max(len(result.clean_df.columns), 1), len(headers) - 1)
    worksheet.set_column("A:A", 24)
    worksheet.set_column("B:B", 18)
    worksheet.set_column("C:C", 28)
    worksheet.set_column("D:D", 12)


def _write_charts(workbook: Any, result: PipelineResult, formats: dict[str, Any]) -> None:
    worksheet = workbook.add_worksheet("Charts")
    category_column = next(
        (
            str(column)
            for column in result.clean_df.select_dtypes(
                include=["object", "string", "category"]
            ).columns
            if 0 < result.clean_df[column].nunique(dropna=False) <= 20
        ),
        None,
    )
    if category_column is not None:
        counts = result.clean_df[category_column].fillna("<missing>").astype(str).value_counts()
        worksheet.write_row(0, 0, [category_column, "Rows"], formats["header"])
        for row_index, (category, count) in enumerate(counts.items(), start=1):
            worksheet.write(row_index, 0, category)
            worksheet.write_number(row_index, 1, int(count), formats["number"])
        category_chart = workbook.add_chart({"type": "column"})
        category_chart.add_series(
            {
                "name": ["Charts", 0, 1],
                "categories": ["Charts", 1, 0, len(counts), 0],
                "values": ["Charts", 1, 1, len(counts), 1],
                "fill": {"color": "#70AD47"},
            }
        )
        category_chart.set_title({"name": f"{category_column} counts"})
        category_chart.set_y_axis({"name": "Rows"})
    else:
        worksheet.write_row(0, 0, ["Metric", "Value"], formats["header"])
        worksheet.write(1, 0, "Rows")
        worksheet.write_number(1, 1, len(result.clean_df), formats["number"])
        category_chart = workbook.add_chart({"type": "column"})
        category_chart.add_series({"categories": "=Charts!$A$2", "values": "=Charts!$B$2"})
        category_chart.set_title({"name": "Dataset size"})

    numeric_columns = [
        str(column)
        for column in result.clean_df.select_dtypes(include="number").columns
        if not pd.api.types.is_bool_dtype(result.clean_df[column])
    ]
    numeric_column = numeric_columns[0] if numeric_columns else None
    second_data_row = 0
    if numeric_column is not None:
        second_data_row = max(len(counts), 1) + 4 if category_column is not None else 5
        worksheet.write_row(second_data_row, 0, ["Row", numeric_column], formats["header"])
        for offset, value in enumerate(result.clean_df[numeric_column].tolist(), start=1):
            worksheet.write_number(second_data_row + offset, 0, offset, formats["number"])
            worksheet.write_number(
                second_data_row + offset,
                1,
                float(value),
                formats["number"],
            )
        line_chart = workbook.add_chart({"type": "line"})
        line_chart.add_series(
            {
                "name": ["Charts", second_data_row, 1],
                "categories": [
                    "Charts",
                    second_data_row + 1,
                    0,
                    second_data_row + len(result.clean_df),
                    0,
                ],
                "values": [
                    "Charts",
                    second_data_row + 1,
                    1,
                    second_data_row + len(result.clean_df),
                    1,
                ],
                "line": {"color": "#ED7D31"},
            }
        )
        line_chart.set_title({"name": f"{numeric_column} by row"})
    else:
        line_chart = workbook.add_chart({"type": "line"})
        line_chart.add_series({"categories": "=Charts!$A$2", "values": "=Charts!$B$2"})
        line_chart.set_title({"name": "Dataset size trend"})

    worksheet.insert_chart("D2", category_chart, {"x_scale": 1.25, "y_scale": 1.15})
    worksheet.insert_chart("D20", line_chart, {"x_scale": 1.25, "y_scale": 1.15})
    worksheet.set_column("A:B", 22)


def build_workbook(result: PipelineResult, out_path: str | Path | None = None) -> Path:
    """Build an auditable six-sheet workbook from a pipeline result."""

    if out_path is None:
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        out_path = Path("outputs/excel") / f"dataset_{timestamp}.xlsx"
    output_path = Path(out_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    import xlsxwriter

    workbook = xlsxwriter.Workbook(str(output_path))
    formats = {
        "title": workbook.add_format({"bold": True, "font_size": 16, "font_color": "#1F4E78"}),
        "header": workbook.add_format(
            {"bold": True, "font_color": "#FFFFFF", "bg_color": "#4472C4", "border": 1}
        ),
        "label": workbook.add_format({"bold": True, "bg_color": "#D9EAF7"}),
        "date": workbook.add_format({"num_format": "yyyy-mm-dd"}),
        "date_time": workbook.add_format({"num_format": "yyyy-mm-dd hh:mm:ss"}),
        "number": workbook.add_format({"num_format": "#,##0.00"}),
        "percent": workbook.add_format({"num_format": "0.00%"}),
    }
    try:
        _write_summary(workbook, result, formats)
        cleaned_sheet = workbook.add_worksheet("Cleaned Data")
        _write_dataframe(cleaned_sheet, result.clean_df, formats)
        outlier_column = next(
            (
                index
                for index, column in enumerate(result.clean_df.columns)
                if str(column) == "is_outlier"
            ),
            None,
        )
        if outlier_column is not None and len(result.clean_df.columns):
            outlier_format = workbook.add_format({"bg_color": "#FFC7CE", "font_color": "#9C0006"})
            outlier_letter = xl_col_to_name(outlier_column)
            cleaned_sheet.conditional_format(
                1,
                0,
                max(len(result.clean_df), 1),
                len(result.clean_df.columns) - 1,
                {
                    "type": "formula",
                    "criteria": f"=${outlier_letter}2=TRUE",
                    "format": outlier_format,
                },
            )
        raw_sheet = workbook.add_worksheet("Raw Data")
        _write_dataframe(raw_sheet, result.raw_df, formats)
        _write_dictionary(workbook, result, formats)
        _write_change_log(workbook, result, formats)
        _write_charts(workbook, result, formats)
    finally:
        workbook.close()

    log("bot", "data_excel_export", path=str(output_path))
    return output_path
