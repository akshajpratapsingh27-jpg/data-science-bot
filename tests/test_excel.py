"""Excel tests - exactly three tests as required."""

from pathlib import Path

from openpyxl import load_workbook

from dsbot.data.excel import build_workbook
from dsbot.data.pipeline import run

CSV_PATH = Path(__file__).parents[1] / "data" / "samples" / "dirty_sales.csv"


def test_workbook_has_all_expected_sheets(tmp_path):
    result = run(CSV_PATH)
    output_path = build_workbook(result, tmp_path / "sales.xlsx")
    workbook = load_workbook(output_path, data_only=False)

    assert set(workbook.sheetnames) == {
        "Summary",
        "Cleaned Data",
        "Raw Data",
        "Data Dictionary",
        "Change Log",
        "Charts",
    }


def test_summary_contains_an_excel_formula(tmp_path):
    result = run(CSV_PATH)
    output_path = build_workbook(result, tmp_path / "sales.xlsx")
    workbook = load_workbook(output_path, data_only=False)
    summary = workbook["Summary"]

    formulas = [
        cell.value
        for row in summary.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    ]
    assert any(value.startswith("=") for value in formulas)


def test_change_log_row_count_matches_pipeline_change_log(tmp_path):
    result = run(CSV_PATH)
    output_path = build_workbook(result, tmp_path / "sales.xlsx")
    workbook = load_workbook(output_path, data_only=False)
    change_log = workbook["Change Log"]

    assert change_log.max_row - 1 == len(result.change_log)
