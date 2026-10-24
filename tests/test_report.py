"""Report tests - exactly three tests as required."""

from pathlib import Path

from docx import Document

from dsbot.data.analysis import analyze
from dsbot.data.pipeline import run
from dsbot.data.report import build_report
from dsbot.llm import FakeLLM

CSV_PATH = Path(__file__).parents[1] / "data" / "samples" / "dirty_sales.csv"


def test_computed_facts_match_direct_pandas_calculations(tmp_path):
    result = run(CSV_PATH)
    facts = analyze(result.clean_df, charts_dir=tmp_path / "charts")
    values = {fact["id"]: fact["value"] for fact in facts}

    assert values["row_count"] == len(result.clean_df)
    assert values["numeric_amount_mean"] == result.clean_df["amount"].mean()
    category_means = result.clean_df.groupby("category")["amount"].mean().to_dict()
    assert values["group_category_amount_mean"] == {
        str(category): float(value) for category, value in category_means.items()
    }


def test_invented_number_is_rejected_and_template_fallback_is_used(tmp_path):
    result = run(CSV_PATH)
    facts = analyze(result.clean_df, charts_dir=tmp_path / "charts")
    llm = FakeLLM(lambda *_: "The result was 123456789.")
    output_path = tmp_path / "report.docx"

    build_report(result, facts, output_path, llm=llm)

    text = "\n".join(paragraph.text for paragraph in Document(output_path).paragraphs)
    assert len(llm.calls) == 2
    assert "123456789" not in text
    assert "computed directly from the cleaned data" in text


def test_docx_contains_expected_headings_and_an_image(tmp_path):
    result = run(CSV_PATH)
    facts = analyze(result.clean_df)
    output_path = tmp_path / "report.docx"

    build_report(result, facts, output_path, llm=FakeLLM({"default": "The dataset is ready."}))

    document = Document(output_path)
    headings = {paragraph.text for paragraph in document.paragraphs}
    assert {
        "Executive Summary",
        "Data Quality",
        "Key Findings",
        "Method",
        "Appendix",
    }.issubset(headings)
    assert len(document.inline_shapes) >= 1
